"""
SkillSync Official Dataset Importer
Reads data/SkillSync_Datasets.xlsx and performs idempotent, deterministic upserts into SQLAlchemy database.
Preserves dataset IDs (USR-*, SKL-*, CRS-*, PRJ-*, REQ-*, MENT-*, QUIZ-*, RES-*, NOTIF-*).
Enforces exactly 3 system roles: ADMIN, MANAGER, EMPLOYEE.
SME is an approved capability of an EMPLOYEE (is_sme=True).
"""
import os
import sys
from datetime import datetime, date
import openpyxl
from sqlalchemy.orm import Session
from app.core.database import SessionLocal, engine, Base
from app.core.security import get_password_hash
from app.models.user import User
from app.models.skill import Skill
from app.models.employee_skill import EmployeeSkill
from app.models.course import Course
from app.models.learning_progress import LearningProgress
from app.models.certification import Certification
from app.models.project import Project
from app.models.project_requirement import ProjectRequirement
from app.models.training import TrainingRequest
from app.models.training_assignment import TrainingAssignment
from app.models.mentorship import MentorshipRecord
from app.models.impact_score import ImpactScore
from app.models.quiz import QuizQuestion, QuizResult
from app.models.notification import Notification
from app.models.market_benchmark import MarketBenchmark
from app.models.sme_expertise import SMEExpertise


DEFAULT_PASSWORD_HASH = get_password_hash("password123")


def parse_date(val):
    if not val:
        return None
    if isinstance(val, (datetime, date)):
        return val if isinstance(val, datetime) else datetime.combine(val, datetime.min.time())
    try:
        s = str(val).strip()
        if " " in s:
            return datetime.strptime(s[:19], "%Y-%m-%d %H:%M:%S")
        return datetime.strptime(s[:10], "%Y-%m-%d")
    except Exception:
        return None


def import_dataset(excel_path: str = "data/SkillSync_Datasets.xlsx", db: Session = None, reset: bool = False):
    close_db = False
    if db is None:
        db = SessionLocal()
        close_db = True

    if reset:
        print("Resetting database schema...")
        Base.metadata.drop_all(bind=engine)

    # Ensure all tables created
    Base.metadata.create_all(bind=engine)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    counts = {}

    print(f"Loading dataset from {excel_path}...")

    # -------------------------------------------------------------
    # 1. Users (50 rows)
    # -------------------------------------------------------------
    ws_users = wb["Users"]
    user_rows = list(ws_users.iter_rows(values_only=True))[1:]
    user_code_map = {}  # "USR-0001" -> User instance

    for r in user_rows:
        if not r or not r[0]:
            continue
        user_code, name, email, raw_role, dept, loc, exp_lvl, exp_yrs, raw_sme, created_at = r[:10]
        
        # Enforce exactly 3 system roles: ADMIN, MANAGER, EMPLOYEE
        # SME is an approved capability of an EMPLOYEE
        raw_role_str = str(raw_role or "").lower()
        if raw_role_str == "manager":
            system_role = "MANAGER"
            is_sme = False
        elif raw_role_str == "sme":
            system_role = "EMPLOYEE"
            is_sme = True
        elif raw_role_str == "admin":
            system_role = "ADMIN"
            is_sme = False
        else:  # learner / employee
            system_role = "EMPLOYEE"
            is_sme = bool(raw_sme)

        existing = db.query(User).filter((User.user_code == user_code) | (User.email == email)).first()
        if not existing:
            user = User(
                user_code=user_code,
                name=name,
                email=email,
                password_hash=DEFAULT_PASSWORD_HASH,
                role=system_role,
                department=dept,
                location=loc,
                experience_level=exp_lvl.capitalize() if exp_lvl else "Junior",
                experience_years=int(exp_yrs or 0),
                is_sme=is_sme,
                is_active=True,
                created_at=parse_date(created_at) or datetime.utcnow(),
            )
            db.add(user)
            db.flush()
            user_code_map[user_code] = user
        else:
            existing.user_code = user_code
            existing.name = name
            existing.role = system_role
            existing.department = dept
            existing.location = loc
            existing.experience_level = exp_lvl.capitalize() if exp_lvl else existing.experience_level
            existing.experience_years = int(exp_yrs or existing.experience_years or 0)
            existing.is_sme = is_sme
            user_code_map[user_code] = existing

    # Provision standard demo personas
    demo_accounts = [
        ("admin@skillsync.com", "Admin User", "ADMIN", "Executive", "New York", "Senior", 10, False, "USR-ADMIN"),
        ("manager@skillsync.com", "Engineering Manager", "MANAGER", "Engineering", "San Francisco", "Senior", 8, False, "USR-MGR"),
        ("sme@skillsync.com", "Senior SME Lead", "EMPLOYEE", "Engineering", "San Francisco", "Senior", 7, True, "USR-SME"),
        ("employee@skillsync.com", "Alex Taylor", "EMPLOYEE", "Engineering", "San Francisco", "Junior", 2, False, "USR-EMP"),
    ]
    mgr_user = None
    for d_email, d_name, d_role, d_dept, d_loc, d_lvl, d_yrs, d_sme, d_code in demo_accounts:
        existing = db.query(User).filter(User.email == d_email).first()
        if not existing:
            u = User(
                user_code=d_code,
                name=d_name,
                email=d_email,
                password_hash=DEFAULT_PASSWORD_HASH,
                role=d_role,
                department=d_dept,
                location=d_loc,
                experience_level=d_lvl,
                experience_years=d_yrs,
                is_sme=d_sme,
                is_active=True,
                created_at=datetime.utcnow(),
            )
            db.add(u)
            db.flush()
            if d_role == "MANAGER":
                mgr_user = u
        else:
            existing.role = d_role
            existing.is_sme = d_sme
            if d_role == "MANAGER":
                mgr_user = existing

    # Assign managers by department
    dept_managers = {}
    for u in db.query(User).filter(User.role == "MANAGER").all():
        if u.department:
            dept_managers[u.department] = u.id

    fallback_mgr_id = mgr_user.id if mgr_user else list(dept_managers.values())[0]
    for u in db.query(User).filter(User.role == "EMPLOYEE").all():
        if not u.manager_id:
            u.manager_id = dept_managers.get(u.department, fallback_mgr_id)

    db.commit()
    counts["users"] = db.query(User).count()
    print(f"  Users: {counts['users']} total")

    # -------------------------------------------------------------
    # 2. Skills_Reference (28 rows)
    # -------------------------------------------------------------
    ws_skills = wb["Skills_Reference"]
    skill_rows = list(ws_skills.iter_rows(values_only=True))[1:]
    skill_code_map = {}
    skill_name_map = {}

    for r in skill_rows:
        if not r or not r[0]:
            continue
        skill_id, skill_name, category, is_active, min_lvl, max_lvl = r[:6]
        existing = db.query(Skill).filter((Skill.skill_code == skill_id) | (Skill.name == skill_name)).first()
        if not existing:
            skill = Skill(
                skill_code=skill_id,
                name=skill_name,
                category=category or "General",
                is_active=bool(is_active),
                min_level=int(min_lvl or 1),
                max_level=int(max_lvl or 5),
            )
            db.add(skill)
            db.flush()
            skill_code_map[skill_id] = skill
            skill_name_map[skill_name.lower()] = skill
        else:
            existing.skill_code = skill_id
            existing.name = skill_name
            existing.category = category or existing.category
            existing.is_active = bool(is_active)
            skill_code_map[skill_id] = existing
            skill_name_map[skill_name.lower()] = existing

    db.commit()
    counts["skills"] = db.query(Skill).count()
    print(f"  Skills: {counts['skills']} total")

    # -------------------------------------------------------------
    # 3. Market_Benchmarks (28 rows)
    # -------------------------------------------------------------
    ws_bench = wb["Market_Benchmarks"]
    bench_rows = list(ws_bench.iter_rows(values_only=True))[1:]
    for r in bench_rows:
        if not r or not r[0]:
            continue
        s_name, cat, g_avg, us_avg, eu_avg, as_avg, growth, d_score, sal_prem, trend, f1, f3 = r[:12]
        existing = db.query(MarketBenchmark).filter(MarketBenchmark.skill_name == s_name).first()
        if not existing:
            mb = MarketBenchmark(
                skill_name=s_name,
                category=cat,
                global_market_avg=g_avg,
                us_market_avg=us_avg,
                europe_market_avg=eu_avg,
                asia_market_avg=as_avg,
                yoy_growth_rate=growth,
                demand_score=d_score,
                salary_premium_percent=sal_prem,
                trend=trend,
                forecast_1yr_demand=f1,
                forecast_3yr_demand=f3,
            )
            db.add(mb)
    db.commit()
    counts["market_benchmarks"] = db.query(MarketBenchmark).count()
    print(f"  Market Benchmarks: {counts['market_benchmarks']} total")

    # -------------------------------------------------------------
    # 4. Skill_Profiles (558 rows)
    # -------------------------------------------------------------
    ws_profiles = wb["Skill_Profiles"]
    profile_rows = list(ws_profiles.iter_rows(values_only=True))[1:]
    for r in profile_rows:
        if not r or not r[0]:
            continue
        (u_id, u_name, dept, loc, exp_lvl, s_id, s_name, cat,
         self_ass, val_lvl, cur_lvl, tgt_lvl, lvl_chg, mkt_avg,
         mkt_pct, abv_mkt, mkt_gap, last_ass, has_cert, has_ment) = r[:20]

        user = user_code_map.get(u_id) or db.query(User).filter(User.user_code == u_id).first()
        skill = skill_code_map.get(s_id) or skill_name_map.get(str(s_name).lower()) or db.query(Skill).filter((Skill.skill_code == s_id) | (Skill.name == s_name)).first()

        if not user or not skill:
            continue

        # Prefer cur_lvl if provided, else round(val_lvl), else round(self_ass or 1)
        effective_level = int(cur_lvl) if cur_lvl is not None else (
            round(val_lvl) if val_lvl is not None else (round(self_ass) if self_ass is not None else 1)
        )
        effective_level = max(1, min(5, effective_level))

        existing = db.query(EmployeeSkill).filter(
            EmployeeSkill.employee_id == user.id,
            EmployeeSkill.skill_id == skill.id
        ).first()

        if not existing:
            es = EmployeeSkill(
                employee_id=user.id,
                skill_id=skill.id,
                current_level=effective_level,
                self_assessment=float(self_ass) if self_ass is not None else None,
                validated_level=float(val_lvl) if val_lvl is not None else None,
                target_level=float(tgt_lvl) if tgt_lvl is not None else None,
                level_change_12mo=float(lvl_chg) if lvl_chg is not None else None,
                market_average=float(mkt_avg) if mkt_avg is not None else None,
                market_percentile=float(mkt_pct) if mkt_pct is not None else None,
                above_market=bool(abv_mkt),
                market_gap=float(mkt_gap) if mkt_gap is not None else None,
                last_assessed=parse_date(last_ass),
                has_certification=bool(has_cert),
                has_mentorship=bool(has_ment),
                validated=bool(val_lvl is not None),
                validated_at=parse_date(last_ass) if val_lvl is not None else None,
            )
            db.add(es)
        else:
            existing.current_level = effective_level
            existing.self_assessment = float(self_ass) if self_ass is not None else existing.self_assessment
            existing.validated_level = float(val_lvl) if val_lvl is not None else existing.validated_level
            existing.target_level = float(tgt_lvl) if tgt_lvl is not None else existing.target_level
            existing.validated = bool(val_lvl is not None)

        # If user is approved SME and holds high proficiency, register in SMEExpertise
        if user.is_sme and effective_level >= 4:
            has_exp = db.query(SMEExpertise).filter(
                SMEExpertise.employee_id == user.id,
                SMEExpertise.skill_id == skill.id
            ).first()
            if not has_exp:
                db.add(SMEExpertise(
                    employee_id=user.id,
                    skill_id=skill.id,
                    skill_level=effective_level,
                    approved_at=datetime.utcnow()
                ))

    db.commit()
    counts["skill_profiles"] = db.query(EmployeeSkill).count()
    print(f"  Skill Profiles: {counts['skill_profiles']} total")

    # -------------------------------------------------------------
    # 5. Courses (15 rows)
    # -------------------------------------------------------------
    ws_courses = wb["Courses"]
    course_rows = list(ws_courses.iter_rows(values_only=True))[1:]
    course_code_map = {}

    for r in course_rows:
        if not r or not r[0]:
            continue
        c_id, title, desc, s_name, s_id, cat, diff, platform, hours, rating, enrolled, comp_rate, prereq, cert_prov, url = r[:15]
        skill = skill_code_map.get(s_id) or skill_name_map.get(str(s_name).lower()) or db.query(Skill).filter((Skill.skill_code == s_id) | (Skill.name == s_name)).first()
        if not skill:
            continue

        diff_str = str(diff)
        if diff_str in ("1", "1.0"):
            diff_str = "Beginner"
        elif diff_str in ("2", "2.0", "3", "3.0"):
            diff_str = "Intermediate"
        elif diff_str in ("4", "4.0", "5", "5.0"):
            diff_str = "Advanced"

        existing = db.query(Course).filter((Course.course_code == c_id) | (Course.title == title)).first()
        if not existing:
            c = Course(
                course_code=c_id,
                title=title,
                description=desc,
                skill_id=skill.id,
                category=cat,
                difficulty=diff_str or "Intermediate",
                platform=platform or "Internal",
                duration_hours=float(hours or 10.0),
                rating=float(rating) if rating is not None else 4.5,
                enrolled_students=int(enrolled or 0) if enrolled is not None else None,
                completion_rate=float(comp_rate) if comp_rate is not None else None,
                prerequisite_skills=prereq,
                certification_provided=bool(cert_prov),
                url=url or f"https://learning.skillsync.com/courses/{c_id}",
            )
            db.add(c)
            db.flush()
            course_code_map[c_id] = c
        else:
            existing.course_code = c_id
            course_code_map[c_id] = existing

    db.commit()
    counts["courses"] = db.query(Course).count()
    print(f"  Courses: {counts['courses']} total")

    # -------------------------------------------------------------
    # 6. Course_Progress (100 rows)
    # -------------------------------------------------------------
    ws_progress = wb["Course_Progress"]
    prog_rows = list(ws_progress.iter_rows(values_only=True))[1:]
    for r in prog_rows:
        if not r or not r[0]:
            continue
        p_id, u_id, u_name, c_id, c_title, pct, status, started, completed, hours, cert_obt = r[:11]
        user = user_code_map.get(u_id) or db.query(User).filter(User.user_code == u_id).first()
        course = course_code_map.get(c_id) or db.query(Course).filter(Course.course_code == c_id).first()
        if not user or not course:
            continue

        existing = db.query(LearningProgress).filter(
            (LearningProgress.progress_code == p_id) |
            ((LearningProgress.employee_id == user.id) & (LearningProgress.course_id == course.id))
        ).first()

        if not existing:
            lp = LearningProgress(
                progress_code=p_id,
                employee_id=user.id,
                course_id=course.id,
                progress_percentage=float(pct or 25.0) if pct is not None else 30.0,
                status=str(status or "IN_PROGRESS").upper(),
                started_at=parse_date(started) or datetime.utcnow(),
                completed_at=parse_date(completed),
                hours_spent=float(hours or 4.0) if hours is not None else 4.0,
                certificate_obtained=bool(cert_obt),
            )
            db.add(lp)

    db.commit()
    counts["course_progress"] = db.query(LearningProgress).count()
    print(f"  Course Progress: {counts['course_progress']} total")

    # -------------------------------------------------------------
    # 7. Certifications (37 rows)
    # -------------------------------------------------------------
    ws_certs = wb["Certifications"]
    cert_rows = list(ws_certs.iter_rows(values_only=True))[1:]
    for r in cert_rows:
        if not r or not r[0]:
            continue
        cert_id, u_id, u_name, s_id, s_name, title, issuer, platform, url, issued, expires, verified = r[:12]
        user = user_code_map.get(u_id) or db.query(User).filter(User.user_code == u_id).first()
        skill = skill_code_map.get(s_id) or skill_name_map.get(str(s_name).lower()) or db.query(Skill).filter((Skill.skill_code == s_id) | (Skill.name == s_name)).first()
        if not user:
            continue

        existing = db.query(Certification).filter(
            (Certification.certification_code == cert_id) |
            ((Certification.user_id == user.id) & (Certification.title == title))
        ).first()

        if not existing:
            cert = Certification(
                certification_code=cert_id,
                user_id=user.id,
                skill_id=skill.id if skill else None,
                skill_name=s_name,
                title=title,
                issuer=issuer or "Accredited Provider",
                platform=platform or "Online",
                url=url,
                issued_at=parse_date(issued),
                expires_at=parse_date(expires),
                verified=bool(verified if verified is not None else 1),
            )
            db.add(cert)

    db.commit()
    counts["certifications"] = db.query(Certification).count()
    print(f"  Certifications: {counts['certifications']} total")

    # -------------------------------------------------------------
    # 8. Project_Requirements (44 rows)
    # -------------------------------------------------------------
    ws_prj = wb["Project_Requirements"]
    prj_rows = list(ws_prj.iter_rows(values_only=True))[1:]
    project_code_map = {}

    for r in prj_rows:
        if not r or not r[0]:
            continue
        p_id, p_name, dept, m_id, m_name, deadline, status, priority, s_name, s_id, req_lvl, headcount, cur_avail, gap, hrs_needed = r[:15]
        
        manager = user_code_map.get(m_id) or db.query(User).filter(User.user_code == m_id).first()
        if not manager:
            manager = db.query(User).filter(User.role == "MANAGER").first()

        skill = skill_code_map.get(s_id) or skill_name_map.get(str(s_name).lower()) or db.query(Skill).filter((Skill.skill_code == s_id) | (Skill.name == s_name)).first()
        if not skill:
            continue

        project = project_code_map.get(p_id) or db.query(Project).filter(Project.project_code == p_id).first()
        if not project:
            project = Project(
                project_code=p_id,
                name=p_name,
                department=dept,
                manager_id=manager.id if manager else fallback_mgr_id,
                deadline=str(deadline or "2026-12-31"),
                status=str(status or "ACTIVE").upper(),
                priority=str(priority or "CRITICAL").upper(),
            )
            db.add(project)
            db.flush()
            project_code_map[p_id] = project

        existing_req = db.query(ProjectRequirement).filter(
            ProjectRequirement.project_id == project.id,
            ProjectRequirement.skill_id == skill.id
        ).first()

        if not existing_req:
            pr = ProjectRequirement(
                project_id=project.id,
                skill_id=skill.id,
                required_people=int(headcount or 1),
                minimum_level=int(req_lvl or 3),
                current_availability=int(cur_avail or 0),
                gap=int(gap or 0),
                training_hours_needed=int(hrs_needed or 0),
            )
            db.add(pr)

    db.commit()
    counts["projects"] = db.query(Project).count()
    counts["project_requirements"] = db.query(ProjectRequirement).count()
    print(f"  Projects: {counts['projects']}, Project Requirements: {counts['project_requirements']} total")

    # -------------------------------------------------------------
    # 9. Training_Requests (40 rows)
    # -------------------------------------------------------------
    ws_train = wb["Training_Requests"]
    train_rows = list(ws_train.iter_rows(values_only=True))[1:]
    for r in train_rows:
        if not r or not r[0]:
            continue
        req_id, req_uid, req_uname, dept, s_name, s_id, req_type, priority, status, sme_uid, sme_uname, req_date, pref_date, notes, pref_sme_uid, pref_sme_uname = r[:16]
        user = user_code_map.get(req_uid) or db.query(User).filter(User.user_code == req_uid).first()
        skill = skill_code_map.get(s_id) or skill_name_map.get(str(s_name).lower()) or db.query(Skill).filter((Skill.skill_code == s_id) | (Skill.name == s_name)).first()
        sme_user = user_code_map.get(sme_uid) or (db.query(User).filter(User.user_code == sme_uid).first() if sme_uid else None)

        if not user or not skill:
            continue

        raw_status = str(status or "pending").upper()
        if raw_status == "OPEN":
            raw_status = "PENDING"

        existing = db.query(TrainingRequest).filter(TrainingRequest.request_code == req_id).first()
        if not existing:
            tr = TrainingRequest(
                request_code=req_id,
                employee_id=user.id,
                skill_id=skill.id,
                requested_level=3,
                request_type=str(req_type or "online_session"),
                priority=str(priority or "MEDIUM").upper(),
                status=raw_status,
                claimed_by_sme_id=sme_user.id if sme_user else None,
                requested_date=str(req_date) if req_date else None,
                preferred_date=str(pref_date) if pref_date else None,
                notes=notes or f"Training request for {s_name}",
                message=notes,
            )
            db.add(tr)
            if sme_user:
                db.flush()
                db.add(TrainingAssignment(
                    training_request_id=tr.id,
                    sme_id=sme_user.id,
                    status="COMPLETED" if raw_status == "COMPLETED" else "ACTIVE"
                ))

    db.commit()
    counts["training_requests"] = db.query(TrainingRequest).count()
    print(f"  Training Requests: {counts['training_requests']} total")

    # -------------------------------------------------------------
    # 10. Mentorship_Records (50 rows)
    # -------------------------------------------------------------
    ws_ment = wb["Mentorship_Records"]
    ment_rows = list(ws_ment.iter_rows(values_only=True))[1:]
    for r in ment_rows:
        if not r or not r[0]:
            continue
        m_id, mentor_uid, mentor_name, mentee_uid, mentee_name, s_name, start_lvl, end_lvl, imp, sess, status, started, completed, rating, impact = r[:15]
        mentor = user_code_map.get(mentor_uid) or db.query(User).filter(User.user_code == mentor_uid).first()
        mentee = user_code_map.get(mentee_uid) or db.query(User).filter(User.user_code == mentee_uid).first()
        if not mentor or not mentee:
            continue

        skill = skill_name_map.get(str(s_name).lower()) or db.query(Skill).filter(Skill.name == s_name).first()

        existing = db.query(MentorshipRecord).filter(MentorshipRecord.mentorship_code == m_id).first()
        if not existing:
            mr = MentorshipRecord(
                mentorship_code=m_id,
                mentor_id=mentor.id,
                mentee_id=mentee.id,
                skill_id=skill.id if skill else None,
                skill_name=s_name,
                start_level=float(start_lvl) if start_lvl is not None else None,
                end_level=float(end_lvl) if end_lvl is not None else None,
                improvement=float(imp) if imp is not None else None,
                sessions_completed=int(sess or 0) if sess is not None else 0,
                status=str(status or "COMPLETED").upper(),
                started_at=parse_date(started),
                completed_at=parse_date(completed),
                mentee_rating=float(rating) if rating is not None else 5.0,
                impact_score=float(impact) if impact is not None else None,
            )
            db.add(mr)

    db.commit()
    counts["mentorship_records"] = db.query(MentorshipRecord).count()
    print(f"  Mentorship Records: {counts['mentorship_records']} total")

    # -------------------------------------------------------------
    # 11. Impact_Scores (3 rows)
    # -------------------------------------------------------------
    ws_imp = wb["Impact_Scores"]
    imp_rows = list(ws_imp.iter_rows(values_only=True))[1:]
    for r in imp_rows:
        if not r or not r[0]:
            continue
        mentor_uid, mentor_name, mentees, lvlups, avg_imp, sessions, imp_score, calc_at = r[:8]
        mentor = user_code_map.get(mentor_uid) or db.query(User).filter(User.user_code == mentor_uid).first()
        if not mentor:
            continue

        existing = db.query(ImpactScore).filter(ImpactScore.mentor_id == mentor.id).first()
        if not existing:
            isc = ImpactScore(
                mentor_id=mentor.id,
                total_mentees=int(mentees or 0),
                level_ups_achieved=int(lvlups or 0),
                avg_improvement=float(avg_imp or 0.0),
                total_sessions=int(sessions or 0),
                impact_score=float(imp_score or 4.5) if imp_score is not None else 4.5,
                last_calculated_at=parse_date(calc_at) or datetime.utcnow(),
            )
            db.add(isc)

    db.commit()
    counts["impact_scores"] = db.query(ImpactScore).count()
    print(f"  Impact Scores: {counts['impact_scores']} total")

    # -------------------------------------------------------------
    # 12. Quiz_Questions (68 rows) & 13. Quiz_Results (100 rows)
    # -------------------------------------------------------------
    ws_quiz_q = wb["Quiz_Questions"]
    quiz_q_rows = list(ws_quiz_q.iter_rows(values_only=True))[1:]
    quiz_code_map = {}
    for r in quiz_q_rows:
        if not r or not r[0]:
            continue
        q_id, s_id, s_name, question, opt_a, opt_b, opt_c, opt_d, ans, points = r[:10]
        skill = skill_code_map.get(s_id) or skill_name_map.get(str(s_name).lower()) or db.query(Skill).filter((Skill.skill_code == s_id) | (Skill.name == s_name)).first()

        existing = db.query(QuizQuestion).filter(QuizQuestion.quiz_code == q_id).first()
        if not existing:
            qq = QuizQuestion(
                quiz_code=q_id,
                skill_id=skill.id if skill else None,
                skill_name=s_name,
                question=question,
                option_a=str(opt_a),
                option_b=str(opt_b),
                option_c=str(opt_c),
                option_d=str(opt_d),
                correct_answer=str(ans).lower() if ans else "a",
                points=int(points or 10),
            )
            db.add(qq)
            db.flush()
            quiz_code_map[q_id] = qq
        else:
            quiz_code_map[q_id] = existing

    db.commit()
    counts["quiz_questions"] = db.query(QuizQuestion).count()
    print(f"  Quiz Questions: {counts['quiz_questions']} total")

    ws_quiz_res = wb["Quiz_Results"]
    quiz_res_rows = list(ws_quiz_res.iter_rows(values_only=True))[1:]
    for r in quiz_res_rows:
        if not r or not r[0]:
            continue
        res_id, u_id, u_name, q_id, s_name, score_pct, pts_earn, total_pts, passed, attempted = r[:10]
        user = user_code_map.get(u_id) or db.query(User).filter(User.user_code == u_id).first()
        quiz = quiz_code_map.get(q_id) or db.query(QuizQuestion).filter(QuizQuestion.quiz_code == q_id).first()
        if not user:
            continue

        existing = db.query(QuizResult).filter(QuizResult.result_code == res_id).first()
        if not existing:
            qr = QuizResult(
                result_code=res_id,
                user_id=user.id,
                quiz_id=quiz.id if quiz else None,
                skill_name=s_name,
                score_percent=float(score_pct or 0.0),
                points_earned=float(pts_earn or 0.0),
                total_points=float(total_pts or 20.0),
                passed=bool(passed),
                attempted_at=parse_date(attempted),
            )
            db.add(qr)

    db.commit()
    counts["quiz_results"] = db.query(QuizResult).count()
    print(f"  Quiz Results: {counts['quiz_results']} total")

    # -------------------------------------------------------------
    # 14. Notifications (150 rows)
    # -------------------------------------------------------------
    ws_notif = wb["Notifications"]
    notif_rows = list(ws_notif.iter_rows(values_only=True))[1:]
    for r in notif_rows:
        if not r or not r[0]:
            continue
        notif_id, u_id, u_name, n_type, msg, is_read, created = r[:7]
        user = user_code_map.get(u_id) or db.query(User).filter(User.user_code == u_id).first()
        if not user:
            continue

        existing = db.query(Notification).filter(Notification.notification_code == notif_id).first()
        if not existing:
            n = Notification(
                notification_code=notif_id,
                user_id=user.id,
                title=str(n_type or "Notice").replace("_", " ").title(),
                message=msg or "",
                type=str(n_type or "system"),
                is_read=bool(is_read),
                created_at=parse_date(created) or datetime.utcnow(),
            )
            db.add(n)

    db.commit()
    counts["notifications"] = db.query(Notification).count()
    print(f"  Notifications: {counts['notifications']} total")

    if close_db:
        db.close()

    print("Dataset import completed successfully!")
    return counts


if __name__ == "__main__":
    import_dataset()
